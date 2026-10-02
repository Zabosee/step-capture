import pytest
from docx import Document
from PIL import Image

from process_recorder.exporter import (_prepare_image, _prepare_zoom, export_docx,
                                       export_document, mark_click, zoom_box)
from process_recorder.models import Step
from process_recorder.pdf_export import export_pdf


@pytest.fixture
def shot(tmp_path):
    """Ein 1920x1200-Screenshot mit weißem Hintergrund."""
    path = tmp_path / "shot.jpg"
    Image.new("RGB", (1920, 1200), "white").save(path)
    return path


def _steps(shot):
    return [
        Step("click", 1, image_path=shot, click_rel=(960, 600), target="Schaltfläche „OK“ – Editor"),
        Step("text", 2, image_path=shot, text="Grüße Ölweg € @\nzweite Zeile"),
        Step("key", 3, image_path=shot, text="Strg+S"),
        Step("protected", 4),
    ]


def test_mark_click_draws_red_only_near_click():
    img = mark_click(Image.new("RGB", (400, 300), "white"), (200, 150))
    r, g, b = img.getpixel((200, 150))
    assert r > 200 and g < 80 and b < 80                        # roter Mittelpunkt
    assert img.getpixel((5, 5)) == (255, 255, 255)              # weit entfernt: unverändert


def test_mark_click_at_edge_does_not_fail():
    for xy in [(0, 0), (399, 299), (-50, 10), (1000, 1000)]:
        mark_click(Image.new("RGB", (400, 300), "white"), xy)


@pytest.mark.parametrize("click", [(0, 0), (1919, 1199), (960, 600), (5, 1190)])
def test_zoom_box_stays_inside_image(click):
    left, top, right, bottom = zoom_box((1920, 1200), click)
    assert 0 <= left < right <= 1920 and 0 <= top < bottom <= 1200
    assert left <= click[0] <= right and top <= click[1] <= bottom


def test_zoom_box_small_image():
    left, top, right, bottom = zoom_box((200, 100), (100, 50))
    assert (left, top) == (0, 0) and (right, bottom) == (200, 100)


def test_prepare_image_is_downscaled_jpeg(shot):
    buf = _prepare_image(Step("click", 0, image_path=shot, click_rel=(100, 100)))
    with Image.open(buf) as img:
        assert img.format == "JPEG" and img.width == 1600


def test_prepare_image_missing_file_returns_none(tmp_path):
    assert _prepare_image(Step("click", 0, image_path=tmp_path / "fehlt.jpg")) is None


def test_prepare_zoom_only_for_clicks(shot):
    assert _prepare_zoom(Step("text", 0, image_path=shot, text="x")) is None
    buf = _prepare_zoom(Step("click", 0, image_path=shot, click_rel=(500, 500)))
    with Image.open(buf) as img:
        assert img.width == 900 and img.height < 900


def test_export_docx_content(tmp_path, shot):
    out = export_docx(_steps(shot), tmp_path / "a.docx", "Monitor 1",
                      title="Meine Anleitung", intro="Einleitung")
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Meine Anleitung" and "Einleitung" in texts
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert headings == ["Schritt 1", "Schritt 2", "Schritt 3", "Schritt 4"]
    assert "Klicken Sie auf die Schaltfläche „OK“." in texts
    assert "Geben Sie den unten stehenden Text ein." in texts
    assert "Drücken Sie die Tastenkombination Strg+S." in texts
    assert "Ausschnitt (vergrößert)" in texts
    assert any("Grüße Ölweg" in cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    # 3 Screenshots + 1 Zoom (nur der Klick)
    assert len(doc.inline_shapes) == 4


def test_export_docx_without_zoom_and_custom_text(tmp_path, shot):
    steps = _steps(shot)
    steps[0].custom = "Auf OK klicken"
    doc = Document(export_docx(steps, tmp_path / "b.docx", zoom=False))
    texts = [p.text for p in doc.paragraphs]
    assert "Auf OK klicken" in texts and "Ausschnitt (vergrößert)" not in texts
    assert len(doc.inline_shapes) == 3


def test_sections_notes_and_outro_are_exported(tmp_path, shot):
    steps = _steps(shot)
    steps[0].section, steps[0].note, steps[0].note_kind = "Anmelden", "Vorher speichern!", "Warnung"
    steps[3].note = "Admin fragen"
    doc = Document(export_docx(steps, tmp_path / "s.docx", outro="Fertig."))
    texts = [p.text for p in doc.paragraphs]
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 1"] == ["Anmelden", "Abschluss"]
    assert texts[-1] == "Fertig."
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "Warnung:\nVorher speichern!" in cells and "Hinweis:\nAdmin fragen" in cells
    pdf = export_pdf(steps, tmp_path / "s.pdf", outro="Fertig.")
    assert pdf.stat().st_size > 1000


def test_export_docx_survives_missing_screenshot(tmp_path, shot):
    steps = _steps(shot)
    steps[0].image_path = tmp_path / "weg.jpg"
    Document(export_docx(steps, tmp_path / "c.docx"))            # darf nicht scheitern


def test_export_pdf_creates_valid_pdf(tmp_path, shot):
    out = export_pdf(_steps(shot), tmp_path / "a.pdf", "Monitor 1", title="PDF-Test")
    data = out.read_bytes()
    assert data.startswith(b"%PDF") and data.rstrip().endswith(b"%%EOF") and len(data) > 5000


def test_export_pdf_without_images(tmp_path):
    steps = [Step("text", 1, text="nur Text"), Step("protected", 2)]
    assert export_pdf(steps, tmp_path / "n.pdf").read_bytes().startswith(b"%PDF")


def test_export_document_dispatches_by_extension(tmp_path, shot):
    steps = _steps(shot)
    assert export_document(steps, tmp_path / "x.docx").read_bytes()[:2] == b"PK"   # zip = docx
    assert export_document(steps, tmp_path / "x.PDF").read_bytes().startswith(b"%PDF")


HOSTILE = [
    Step("click", 1, click_rel=(5, 5), target="Schaltfläche „a\x00b\x08c“ – Fenster \x1f"),
    Step("text", 2, text="pw\x0bx\ud800y"),
    Step("click", 3, click_rel=(5, 5), custom="x\x1fy"),
    Step("text", 4, text="A" * 20000),
    Step("text", 5, text="\n".join(f"zeile {i}" for i in range(3000))),
    Step("text", 6, text="<b>x</b> <img src='/etc/passwd'/> &#0; <font size=99999>"),
]


@pytest.mark.parametrize("ext", ["docx", "pdf"])
def test_export_survives_hostile_text(tmp_path, ext):
    """Steuerzeichen/Surrogate aus fremden Fenstern oder extrem lange Eingaben dürfen den Export nicht abbrechen."""
    out = export_document(HOSTILE, tmp_path / f"x.{ext}", "Mon\x00itor", title="T\x00", intro="i\x01")
    assert out.stat().st_size > 0
    if ext == "docx":
        text = "\n".join(p.text for p in Document(out).paragraphs)
        assert "\x00" not in text
