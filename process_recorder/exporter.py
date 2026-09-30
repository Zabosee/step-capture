"""Word-Export: Anleitung mit markierten Screenshots und kopierbarem Text."""
from __future__ import annotations

import io
import os
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw

from .models import Step

MAX_IMAGE_WIDTH_PX = 1600   # Screenshots werden für die Datei verkleinert
IMAGE_WIDTH_CM = 16.0


def mark_click(img: Image.Image, xy: Tuple[int, int]) -> Image.Image:
    """Zeichnet einen roten Kreis mit Mittelpunkt um die Klickposition (RGB-Bild wird verändert)."""
    x, y = xy
    r = max(22, int(img.width * 0.013))
    w = max(4, r // 5)
    pad = r + w + 2
    # Nur den Ausschnitt um den Klick bearbeiten statt das ganze Bild zu compositen
    box = (max(0, x - pad), max(0, y - pad), min(img.width, x + pad), min(img.height, y + pad))
    if box[0] >= box[2] or box[1] >= box[3]:
        return img
    crop = img.crop(box).convert("RGBA")
    overlay = Image.new("RGBA", crop.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    cx, cy = x - box[0], y - box[1]
    draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=(255, 0, 0, 55),
                 outline=(255, 0, 0, 255), width=w)
    d = max(3, w)
    draw.ellipse((cx - d, cy - d, cx + d, cy + d), fill=(255, 0, 0, 255))
    img.paste(Image.alpha_composite(crop, overlay).convert("RGB"), box[:2])
    return img


def _prepare_image(step: Step) -> Optional[io.BytesIO]:
    if not step.image_path or not Path(step.image_path).exists():
        return None
    with Image.open(step.image_path) as raw:
        img = raw.convert("RGB")
    click = step.click_rel
    if img.width > MAX_IMAGE_WIDTH_PX:                # erst verkleinern (schneller), dann markieren
        factor = MAX_IMAGE_WIDTH_PX / img.width
        img = img.resize((MAX_IMAGE_WIDTH_PX, round(img.height * factor)), Image.BILINEAR)
        if click:
            click = (round(click[0] * factor), round(click[1] * factor))
    if click:
        img = mark_click(img, click)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    buf.seek(0)
    return buf


def _shade_cell(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _add_note_box(doc, label: str, text: str, mono: bool, fill: str) -> None:
    """Einzeilige Tabelle als 'Textbox' mit Beschriftung und Inhalt."""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    _shade_cell(cell, fill)
    p = cell.paragraphs[0]
    run = p.add_run(label)
    run.bold = True
    for line in (text.split("\n") if text else []):
        para = cell.add_paragraph()
        r = para.add_run(line)
        if mono:
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt(11)


def export_docx(steps: Iterable[Step], output: Path, monitor_label: str = "") -> Path:
    """Erzeugt die Word-Anleitung und gibt den Pfad zurück."""
    steps = list(steps)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    doc.add_heading("Prozessdokumentation", level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Erstellt am {datetime.now():%d.%m.%Y um %H:%M} Uhr").italic = True
    if monitor_label:
        meta.add_run(f"\nAufgenommener Bildschirm: {monitor_label}").italic = True
    meta.add_run(f"\nAnzahl Schritte: {len(steps)}").italic = True

    # Bilder parallel vorbereiten (Pillow gibt die GIL frei); Reihenfolge bleibt erhalten
    with ThreadPoolExecutor(max_workers=min(8, os.cpu_count() or 4)) as pool:
        images = list(pool.map(lambda st: None if st.kind == "protected" else _prepare_image(st),
                               steps))

    for number, (step, image) in enumerate(zip(steps, images), start=1):
        heading = doc.add_heading(f"Schritt {number}", level=2)
        heading.paragraph_format.keep_with_next = True
        desc = doc.add_paragraph(f"Schritt {number}: {step.description()}")
        desc.paragraph_format.keep_with_next = True

        if step.kind == "protected":
            _add_note_box(doc, "Hinweis: Dieser Bereich wurde bewusst nicht aufgezeichnet.",
                          "", mono=False, fill="FFF2CC")
            continue

        if image is not None:
            pic_par = doc.add_paragraph()
            pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic_par.paragraph_format.keep_with_next = step.kind == "text"
            pic_par.add_run().add_picture(image, width=Cm(IMAGE_WIDTH_CM))

        if step.text:
            _add_note_box(doc, "Eingegebener Text zum Kopieren:", step.text,
                          mono=True, fill="F2F2F2")
        doc.add_paragraph()

    output = Path(output)
    doc.save(output)
    return output
